"""Synthetic OOXML fixtures, generated locally; no private docs or bundled fonts."""
from __future__ import annotations

import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Callable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree as ET
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor as Color
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches as SlideInches, Pt as SlidePt

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
WP = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS = {'a': A, 'p': P, 'w': W, 'wp': WP, 'r': R}


def edit_zip(source: Path, destination: Path, edit: Callable[[dict[str, bytes]], None]) -> None:
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    edit(parts)
    with zipfile.ZipFile(destination, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def xml_edit(parts: dict[str, bytes], path: str, edit: Callable) -> None:
    root = ET.fromstring(parts[path])
    edit(root)
    parts[path] = ET.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def generate(directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    image = Image.new('RGB', (400, 200), 'red')
    ImageDraw.Draw(image).rectangle((200, 0, 399, 199), fill='blue')
    image.save(directory / 'two_colors.png')
    deck = Presentation()
    deck.slide_width, deck.slide_height = SlideInches(10), SlideInches(5.625)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = Color.from_string('16324F')
    title = slide.shapes.add_textbox(SlideInches(.5), SlideInches(.3), SlideInches(9), SlideInches(.7))
    paragraph = title.text_frame.paragraphs[0]
    paragraph.text = 'First slide / Первая страница'
    paragraph.font.size, paragraph.font.color.rgb = SlidePt(30), Color(255, 255, 255)
    picture = slide.shapes.add_picture(str(directory / 'two_colors.png'), SlideInches(.5), SlideInches(1.5), SlideInches(3), SlideInches(2))
    picture.crop_left = .5
    group = slide.shapes.add_group_shape()
    for x, rgb in [(1, (0, 255, 0)), (2, (255, 255, 0))]:
        shape = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, SlideInches(x), SlideInches(1), SlideInches(1), SlideInches(1))
        shape.fill.solid(); shape.fill.fore_color.rgb = Color(*rgb); shape.line.fill.background()
    group.left, group.top, group.width, group.height = SlideInches(5), SlideInches(1.5), SlideInches(4), SlideInches(2)
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_textbox(SlideInches(1), SlideInches(1), SlideInches(8), SlideInches(1)).text = 'Second slide with default white background'
    second.shapes.add_picture(str(directory / 'two_colors.png'), SlideInches(1), SlideInches(2), SlideInches(2), SlideInches(1))
    deck.save(directory / 'basic.pptx')
    deck.slides._sldIdLst.insert(0, deck.slides._sldIdLst[-1])
    deck.save(directory / 'reordered.pptx')

    def renamed(parts):
        parts['ppt/slides/intro.xml'] = parts.pop('ppt/slides/slide1.xml')
        parts['ppt/slides/_rels/intro.xml.rels'] = parts.pop('ppt/slides/_rels/slide1.xml.rels')
        for path in ['ppt/_rels/presentation.xml.rels', '[Content_Types].xml']:
            parts[path] = parts[path].replace(b'slides/slide1.xml', b'slides/intro.xml')
    edit_zip(directory/'basic.pptx', directory/'renamed.pptx', renamed)
    edit_zip(directory/'basic.pptx', directory/'no_app.pptx', lambda parts: parts.pop('docProps/app.xml'))
    def missing(parts):
        parts.pop('ppt/slides/slide1.xml')
    edit_zip(directory/'basic.pptx', directory/'broken.pptx', missing)

    def bgref(parts):
        def slide_xml(root):
            common = root.find(f'{{{P}}}cSld'); old = common.find(f'{{{P}}}bg'); common.remove(old)
            bg = ET.Element(f'{{{P}}}bg'); ref = ET.SubElement(bg, f'{{{P}}}bgRef', idx='1001')
            ET.SubElement(ref, f'{{{A}}}srgbClr', val='DC7A21'); common.insert(0, bg)
        def theme(root):
            styles = root.find(f'.//{{{A}}}bgFillStyleLst')
            for node in list(styles): styles.remove(node)
            fill = ET.SubElement(styles, f'{{{A}}}solidFill'); ET.SubElement(fill, f'{{{A}}}schemeClr', val='phClr')
        xml_edit(parts, 'ppt/slides/slide1.xml', slide_xml)
        xml_edit(parts, 'ppt/theme/theme1.xml', theme)
    edit_zip(directory/'basic.pptx', directory/'bgref.pptx', bgref)

    def gradient(parts):
        def slide_xml(root):
            bg = root.find(f'.//{{{P}}}bgPr')
            for node in list(bg): bg.remove(node)
            grad = ET.SubElement(bg, f'{{{A}}}gradFill'); stops = ET.SubElement(grad, f'{{{A}}}gsLst')
            for position, colour in [('0', 'FF0000'), ('100000', '0000FF')]:
                stop = ET.SubElement(stops, f'{{{A}}}gs', pos=position); ET.SubElement(stop, f'{{{A}}}srgbClr', val=colour)
            ET.SubElement(grad, f'{{{A}}}lin', ang='0', scaled='1')
        xml_edit(parts, 'ppt/slides/slide1.xml', slide_xml)
    edit_zip(directory/'basic.pptx', directory/'gradient.pptx', gradient)

    def background_image(parts):
        def slide_xml(root):
            bg = root.find(f'.//{{{P}}}bgPr')
            for node in list(bg): bg.remove(node)
            fill = ET.SubElement(bg, f'{{{A}}}blipFill')
            ref = ET.SubElement(fill, f'{{{A}}}blip', {f'{{{R}}}embed': 'rIdImageBackground'})
            ET.SubElement(ref, f'{{{A}}}alphaModFix', amt='50000')
            ET.SubElement(ET.SubElement(fill, f'{{{A}}}stretch'), f'{{{A}}}fillRect')
        xml_edit(parts, 'ppt/slides/slide1.xml', slide_xml)
        def rels(root): ET.SubElement(root, f'{{{REL}}}Relationship', Id='rIdImageBackground', Type=R+'/image', Target='../media/image1.png')
        xml_edit(parts, 'ppt/slides/_rels/slide1.xml.rels', rels)
    edit_zip(directory/'basic.pptx', directory/'background_image.pptx', background_image)

    def text_properties(parts):
        def slide_xml(root):
            para = root.find(f'.//{{{P}}}sp/{{{P}}}txBody/{{{A}}}p')
            default = para.find(f'{{{A}}}pPr/{{{A}}}defRPr'); default.set('b','1'); default.set('kern','1200')
            ET.SubElement(default, f'{{{A}}}latin', typeface='+mn-lt')
            run = para.find(f'{{{A}}}r'); props = ET.Element(f'{{{A}}}rPr', b='0', i='1'); run.insert(0, props)
            body = root.find(f'.//{{{P}}}sp/{{{P}}}txBody/{{{A}}}bodyPr')
            for element in list(body): body.remove(element)
            ET.SubElement(body, f'{{{A}}}normAutofit', fontScale='80000')
        xml_edit(parts, 'ppt/slides/slide1.xml', slide_xml)
    edit_zip(directory/'basic.pptx', directory/'text_properties.pptx', text_properties)

    doc = Document(); section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.header.paragraphs[0].text = 'Header – Колонтитул'
    section.footer.paragraphs[0].text = 'Footer'
    bg = OxmlElement('w:background'); bg.set(qn('w:color'), 'FFF4CC'); doc._element.insert(0, bg)
    doc.add_heading('Word rendering', 0)
    doc.add_paragraph('Привет мир. A page must fit mobile width and preserve its background.')
    doc.add_picture(str(directory/'two_colors.png'), width=Inches(4))
    for i in range(45): doc.add_paragraph(f'Paragraph-{i+1:02d}: Text with a stable line spacing. Текст на русском языке. '*2)
    doc.save(directory/'basic.docx')

    def crop(parts):
        def doc_xml(root):
            fill=root.find(f'.//{{{A}}}blip/..')
            ET.SubElement(fill, f'{{{A}}}srcRect', l='50000', r='0', t='0', b='0')
        xml_edit(parts,'word/document.xml',doc_xml)
    edit_zip(directory/'basic.docx',directory/'crop.docx',crop)
    def flip(parts):
        xml_edit(parts,'word/document.xml', lambda root: root.find(f'.//{{{A}}}xfrm').set('flipH','1'))
    edit_zip(directory/'basic.docx',directory/'flip.docx',flip)

    def anchor(parts):
        def doc_xml(root):
            node = root.find(f'.//{{{WP}}}inline'); node.tag=f'{{{WP}}}anchor'
            node.set('behindDoc','1'); node.set('simplePos','0')
            for name, value in [('positionH','914400'),('positionV','1828800')]:
                axis=ET.SubElement(node,f'{{{WP}}}{name}',relativeFrom='page');ET.SubElement(axis,f'{{{WP}}}posOffset').text=value
            ET.SubElement(node,f'{{{WP}}}wrapNone')
        xml_edit(parts,'word/document.xml',doc_xml)
    edit_zip(directory/'basic.docx',directory/'anchor.docx',anchor)

    table_doc = Document(); table_doc.sections[0].page_height=Inches(6)
    table = table_doc.add_table(rows=70, cols=2); table.style='Table Grid'
    for i,row in enumerate(table.rows):
        row.cells[0].text=f'Row-{i:02d}';row.cells[1].text='Value'
    table.cell(8,0).merge(table.cell(11,0))
    table_doc.save(directory/'table.docx')

    columns = Document(); cols=columns.sections[0]._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'2')
    for i in range(20):columns.add_paragraph(f'Column text {i} '*8)
    columns.save(directory/'columns.docx')
    columns.add_section(WD_SECTION.CONTINUOUS)
    columns.sections[-1]._sectPr.find(qn('w:cols')).set(qn('w:num'),'1')
    columns.add_paragraph('Continuous section preserved')
    columns.save(directory/'continuous.docx')

    breaks=Document();p=breaks.add_paragraph('Before cached break'); r=p.add_run()._r;r.append(OxmlElement('w:lastRenderedPageBreak'));p.add_run('After cached break')
    breaks.save(directory/'break.docx')
    styles=Document();derived=styles.styles.add_style('FolioDerived',WD_STYLE_TYPE.PARAGRAPH);middle=styles.styles.add_style('FolioMiddle',WD_STYLE_TYPE.PARAGRAPH);base=styles.styles.add_style('FolioBase',WD_STYLE_TYPE.PARAGRAPH)
    base.font.size=Pt(22);base.font.color.rgb=RGBColor.from_string('2468AC');middle.base_style=base;derived.base_style=middle
    styles.add_paragraph('Inherited formatting',style=derived);styles.save(directory/'styles.docx')
    # Additional scene fixtures exercise shared master paint servers and transforms.
    def gradient_shapes(parts):
        def add(root):
            tree=root.find(f'.//{{{P}}}spTree')
            source=next(root.iter(f'{{{P}}}sp'))
            for index,colours in enumerate([('FF0000','00FF00'),('0000FF','FFFF00')]):
                shape=deepcopy(source)
                for placeholder in list(shape.iter(f'{{{P}}}ph')):placeholder.getparent().remove(placeholder)
                shape.find(f'{{{P}}}nvSpPr/{{{P}}}cNvPr').set('id',str(100+index))
                props=shape.find(f'{{{P}}}spPr')
                for value in list(props):props.remove(value)
                xfrm=ET.SubElement(props,f'{{{A}}}xfrm')
                ET.SubElement(xfrm,f'{{{A}}}off',x=str((index+1)*914400),y='3657600')
                ET.SubElement(xfrm,f'{{{A}}}ext',cx='914400',cy='914400')
                ET.SubElement(ET.SubElement(props,f'{{{A}}}prstGeom',prst='rect'),f'{{{A}}}avLst')
                gradient=ET.SubElement(props,f'{{{A}}}gradFill');stops=ET.SubElement(gradient,f'{{{A}}}gsLst')
                for pos,colour in zip(['0','20000','100000'],[colours[0],'FFFFFF',colours[1]]):
                    ET.SubElement(ET.SubElement(stops,f'{{{A}}}gs',pos=pos),f'{{{A}}}srgbClr',val=colour)
                ET.SubElement(gradient,f'{{{A}}}lin',ang='0',scaled='1')
                body=shape.find(f'{{{P}}}txBody')
                if body is not None:shape.remove(body)
                tree.append(shape)
        xml_edit(parts,'ppt/slideMasters/slideMaster1.xml',add)
    edit_zip(directory/'basic.pptx',directory/'gradient_shapes.pptx',gradient_shapes)

    def rotated(parts):
        def edit(root):
            xfrm=root.find(f'.//{{{P}}}grpSp/{{{P}}}grpSpPr/{{{A}}}xfrm');xfrm.set('rot','5400000');xfrm.set('flipH','1')
            picture=root.find(f'.//{{{P}}}pic/{{{P}}}spPr/{{{A}}}xfrm');picture.set('flipV','1')
        xml_edit(parts,'ppt/slides/slide1.xml',edit)
    edit_zip(directory/'basic.pptx',directory/'rotated.pptx',rotated)

    def wrapped(parts):
        def edit(root):
            shape=root.find(f'.//{{{P}}}sp');shape.find(f'{{{P}}}spPr/{{{A}}}xfrm/{{{A}}}ext').set('cx','1828800')
            shape.find(f'{{{P}}}spPr/{{{A}}}xfrm/{{{A}}}ext').set('cy','1828800')
            shape.find(f'.//{{{A}}}t').text='One two three four five six seven eight nine ten'
            shape.find(f'.//{{{A}}}defRPr').set('sz','2400')
            body=shape.find(f'{{{P}}}txBody/{{{A}}}bodyPr');body.set('wrap','square');body.set('lIns','182880');body.set('rIns','182880')
        xml_edit(parts,'ppt/slides/slide1.xml',edit)
    edit_zip(directory/'basic.pptx',directory/'wrapped.pptx',wrapped)

    slide_table=Presentation();slide=slide_table.slides.add_slide(slide_table.slide_layouts[6]);table=slide.shapes.add_table(3,2,SlideInches(1),SlideInches(1),SlideInches(7),SlideInches(3)).table
    for row in range(3):
        for column in range(2):
            table.cell(row,column).text=f'Cell {row},{column}'
            table.cell(row,column).fill.solid();table.cell(row,column).fill.fore_color.rgb=Color.from_string('DDDD88')
    slide_table.save(directory/'table.pptx')
